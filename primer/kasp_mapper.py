#!/usr/bin/env python3
"""KASP/SSR primer mapping and reporting for Format1/Format2/SSR TSV inputs."""

from __future__ import annotations

import csv
import re
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Annotated, Any

import typer
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn
from rich.table import Table

# Constants
FAM_TAG = "GAAGGGGACCAAGTTAATGCT"
HEX_TAG = "GAAGCCCGAAGTCAACGGATT"
SNP_PATTERN = re.compile(r"\[([ACGT]+)/([ACGT]+)\]", re.IGNORECASE)
CHROM_SPLIT_PATTERN = re.compile(r"[,;|\s]+")

# Configure Loguru
logger.remove()
logger.add(
    sys.stderr,
    format="<green>{time:HH:mm:ss}</green> | <level>{level: <8}</level> | <level>{message}</level>",
    level="INFO",
    colorize=True,
)

# Rich console
console = Console()

# Typer app
app = typer.Typer(
    name="kasp-mapper",
    help="KASP/SSR primer mapping and reporting for Format1/Format2/SSR TSV inputs",
    rich_markup_mode="rich",
)


@dataclass(frozen=True, slots=True)
class BlastResult:
    query_id: str
    subject_id: str
    identity: float
    alignment_length: int
    mismatches: int
    gap_opens: int
    query_start: int
    query_end: int
    subject_start: int
    subject_end: int
    evalue: float
    bit_score: float


@dataclass(slots=True)
class KaspPrimer:
    primer_id: str
    format_type: str
    fam_primer: str | None = None
    hex_primer: str | None = None
    common_primer: str | None = None
    flank_seq: str | None = None
    snp_pos: int | None = None
    ref_allele: str | None = None
    alt_allele: str | None = None
    forward_primer: str | None = None
    reverse_primer: str | None = None


@dataclass(frozen=True, slots=True)
class ValidSsrLocus:
    chrom: str
    forward_pos: tuple[int, int]
    reverse_pos: tuple[int, int]
    forward_strand: str
    reverse_strand: str
    amplicon_size: int
    forward_identity: float
    reverse_identity: float
    forward_coverage: float
    reverse_coverage: float


@dataclass(frozen=True, slots=True)
class ValidKaspLocus:
    chrom: str
    snp_pos: int
    ref_allele: str
    alt_allele: str
    fam_pos: tuple[int, int]
    hex_pos: tuple[int, int]
    common_pos: tuple[int, int]
    amplicon_size: int
    query_length: int
    alignment_length: int
    coverage: float
    fam_coverage: float
    hex_coverage: float
    common_coverage: float
    fam_identity: float
    hex_identity: float
    common_identity: float


_VALID_DNA_CHARS = set("ACGTN")


def _is_valid_dna(seq: str) -> bool:
    """Return True if seq is non-empty and contains only ACGTN."""
    return bool(seq) and all(ch in _VALID_DNA_CHARS for ch in seq)


def detect_format(file_path: Path) -> str:
    """Auto-detect input format: 'format1', 'format2' or 'ssr'."""
    if not file_path.exists():
        raise FileNotFoundError(f"Input file not found: {file_path}")

    format1_score = 0
    format2_score = 0
    ssr_score = 0

    with file_path.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                continue
            parts = line.split("\t")

            if len(parts) >= 4:
                fam = parts[1].upper()
                hex_seq = parts[2].upper()
                if FAM_TAG in fam or HEX_TAG in hex_seq:
                    format1_score += 2
                else:
                    format1_score += 1

            if len(parts) >= 2 and SNP_PATTERN.search(parts[1]):
                format2_score += 2

            if len(parts) == 3:
                forward = parts[1].strip().upper()
                reverse = parts[2].strip().upper()
                if (
                    _is_valid_dna(forward)
                    and _is_valid_dna(reverse)
                    and FAM_TAG not in forward
                    and HEX_TAG not in reverse
                    and not SNP_PATTERN.search(parts[1])
                    and not SNP_PATTERN.search(parts[2])
                ):
                    ssr_score += 2

    if format1_score == 0 and format2_score == 0 and ssr_score == 0:
        raise ValueError(
            "Cannot detect format: no valid Format1, Format2 or SSR pattern found."
        )

    best = max(
        ("format1", format1_score),
        ("format2", format2_score),
        ("ssr", ssr_score),
        key=lambda item: item[1],
    )
    return best[0]


def parse_kasp1(file_path: Path) -> list[KaspPrimer]:
    """Parse Format1 TSV (ID, FAM, HEX, Common) and remove dye tags."""
    primers: list[KaspPrimer] = []

    if not file_path.exists():
        raise FileNotFoundError(f"Format1 file not found: {file_path}")

    with file_path.open("r", encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        for row in reader:
            if not row or len(row) < 4:
                continue
            primer_id = row[0].strip()
            fam_seq = row[1].strip().upper()
            hex_seq = row[2].strip().upper()
            common_seq = row[3].strip().upper()

            if primer_id.lower() in {"id", "primer_id"}:
                continue

            if fam_seq.startswith(FAM_TAG):
                fam_seq = fam_seq[len(FAM_TAG) :]
            if hex_seq.startswith(HEX_TAG):
                hex_seq = hex_seq[len(HEX_TAG) :]

            primers.append(
                KaspPrimer(
                    primer_id=primer_id,
                    format_type="format1",
                    fam_primer=fam_seq,
                    hex_primer=hex_seq,
                    common_primer=common_seq,
                )
            )

    if not primers:
        raise ValueError("No valid Format1 records were parsed.")
    logger.info(f"Parsed {len(primers)} Format1 primers")
    return primers


def parse_kasp2(file_path: Path) -> list[KaspPrimer]:
    """Parse Format2 TSV (ID, flank with [REF/ALT] marker)."""
    primers: list[KaspPrimer] = []

    if not file_path.exists():
        raise FileNotFoundError(f"Format2 file not found: {file_path}")

    with file_path.open("r", encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        for row in reader:
            if not row or len(row) < 2:
                continue
            primer_id = row[0].strip()
            flank = row[1].strip().upper()

            if primer_id.lower() in {"id", "primer_id"}:
                continue

            match = SNP_PATTERN.search(flank)
            if not match:
                continue

            ref = match.group(1).upper()
            alt = match.group(2).upper()
            marker_start = match.start()
            marker_end = match.end()

            clean_flank = flank[:marker_start] + ref + flank[marker_end:]
            snp_pos_1based = marker_start + 1

            primers.append(
                KaspPrimer(
                    primer_id=primer_id,
                    format_type="format2",
                    flank_seq=clean_flank,
                    snp_pos=snp_pos_1based,
                    ref_allele=ref,
                    alt_allele=alt,
                )
            )

    if not primers:
        raise ValueError("No valid Format2 records were parsed.")
    logger.info(f"Parsed {len(primers)} Format2 primers")
    return primers


def parse_ssr(file_path: Path) -> list[KaspPrimer]:
    """Parse SSR TSV (ID, Forward, Reverse) with ACGTN-only sequences."""
    primers: list[KaspPrimer] = []

    if not file_path.exists():
        raise FileNotFoundError(f"SSR file not found: {file_path}")

    with file_path.open("r", encoding="utf-8") as handle:
        reader = csv.reader(handle, delimiter="\t")
        for row in reader:
            if not row or len(row) < 3:
                continue
            primer_id = row[0].strip()
            forward_seq = row[1].strip().upper()
            reverse_seq = row[2].strip().upper()

            if primer_id.lower() in {"id", "primer_id", "name"}:
                continue

            if not forward_seq or not reverse_seq:
                logger.warning(f"Skip SSR record {primer_id}: empty sequence")
                continue

            if not (_is_valid_dna(forward_seq) and _is_valid_dna(reverse_seq)):
                logger.warning(
                    f"Skip SSR record {primer_id}: sequence contains non-ACGTN characters"
                )
                continue

            primers.append(
                KaspPrimer(
                    primer_id=primer_id,
                    format_type="ssr",
                    forward_primer=forward_seq,
                    reverse_primer=reverse_seq,
                )
            )

    if not primers:
        raise ValueError("No valid SSR records were parsed.")
    logger.info(f"Parsed {len(primers)} SSR primers")
    return primers


def load_id_chr_map(map_file: Path | None) -> dict[str, set[str]]:
    """Load optional primer ID to target chromosome mapping."""
    if map_file is None:
        return {}
    if not map_file.exists():
        raise FileNotFoundError(f"ID-chr map file not found: {map_file}")

    mapping: dict[str, set[str]] = {}
    with map_file.open("r", encoding="utf-8") as handle:
        for line_no, raw_line in enumerate(handle, start=1):
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue

            if "\t" in line:
                parts = line.split("\t")
            elif "," in line:
                parts = line.split(",")
            else:
                parts = line.split()

            if len(parts) < 2:
                raise ValueError(f"Invalid ID-chr map line {line_no}: {raw_line.rstrip()}")

            primer_id = parts[0].strip()
            chrom_field = parts[1].strip()
            if (
                primer_id.lower() in {"id", "primer_id", "marker_id"}
                and chrom_field.lower() in {"chr", "chrom", "chromosome", "subject_id"}
            ):
                continue

            chroms = {chrom for chrom in CHROM_SPLIT_PATTERN.split(chrom_field) if chrom}
            if not primer_id or not chroms:
                raise ValueError(f"Invalid ID-chr map line {line_no}: {raw_line.rstrip()}")

            mapping.setdefault(primer_id, set()).update(chroms)

    logger.info(f"Loaded chromosome hints for {len(mapping)} primer IDs")
    return mapping


def create_blast_query_file(primers: list[KaspPrimer], output_file: Path) -> None:
    """Create BLAST query FASTA from parsed KASP/SSR primer records."""
    with output_file.open("w", encoding="utf-8") as handle:
        for primer in primers:
            if primer.format_type == "format1":
                if primer.fam_primer:
                    handle.write(f">{primer.primer_id}_FAM\n{primer.fam_primer}\n")
                if primer.hex_primer:
                    handle.write(f">{primer.primer_id}_HEX\n{primer.hex_primer}\n")
                if primer.common_primer:
                    handle.write(f">{primer.primer_id}_Common\n{primer.common_primer}\n")
            elif primer.format_type == "ssr":
                if primer.forward_primer:
                    handle.write(f">{primer.primer_id}_F\n{primer.forward_primer}\n")
                if primer.reverse_primer:
                    handle.write(f">{primer.primer_id}_R\n{primer.reverse_primer}\n")
            elif primer.flank_seq:
                handle.write(f">{primer.primer_id}\n{primer.flank_seq}\n")

    logger.debug(f"Created BLAST query file: {output_file}")


def run_blast(
    query_file: Path,
    db_path: Path,
    output_file: Path,
    evalue: float = 1000,
    word_size: int = 11,
    threads: int = 1,
) -> None:
    """Run blastn-short and write tabular output with qlen."""
    cmd = [
        "blastn",
        "-query", str(query_file),
        "-db", str(db_path),
        "-out", str(output_file),
        "-task", "blastn-short",
        "-evalue", str(evalue),
        "-word_size", str(word_size),
        "-num_threads", str(threads),
        "-dust", "no",
        "-outfmt", "6 std qlen",
    ]

    logger.debug(f"Running BLAST: {' '.join(cmd)}")
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        stderr = result.stderr.strip() or "Unknown BLAST error"
        raise RuntimeError(f"BLAST failed: {stderr}")
    logger.debug(f"BLAST completed: {output_file}")


def parse_blast_results(blast_output: Path) -> list[BlastResult]:
    """Parse tabular BLAST output (outfmt 6 std qlen)."""
    results: list[BlastResult] = []

    if not blast_output.exists():
        raise FileNotFoundError(f"BLAST output not found: {blast_output}")

    with blast_output.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) < 12:
                continue
            results.append(
                BlastResult(
                    query_id=parts[0],
                    subject_id=parts[1],
                    identity=float(parts[2]),
                    alignment_length=int(parts[3]),
                    mismatches=int(parts[4]),
                    gap_opens=int(parts[5]),
                    query_start=int(parts[6]),
                    query_end=int(parts[7]),
                    subject_start=int(parts[8]),
                    subject_end=int(parts[9]),
                    evalue=float(parts[10]),
                    bit_score=float(parts[11]),
                )
            )
    logger.debug(f"Parsed {len(results)} BLAST results")
    return results


def _query_to_subject_pos(hit: BlastResult, query_pos_1based: int) -> int | None:
    """Map a 1-based query coordinate to 1-based subject coordinate for one hit."""
    qmin = min(hit.query_start, hit.query_end)
    qmax = max(hit.query_start, hit.query_end)
    if query_pos_1based < qmin or query_pos_1based > qmax:
        return None

    if hit.query_end >= hit.query_start:
        q_offset = query_pos_1based - hit.query_start
    else:
        q_offset = hit.query_start - query_pos_1based

    if hit.subject_end >= hit.subject_start:
        return hit.subject_start + q_offset
    return hit.subject_start - q_offset


def _format_interval(start: int, end: int) -> tuple[int, int]:
    """Return normalized genomic interval as (min, max)."""
    return (min(start, end), max(start, end))


def _query_coverage(hit: BlastResult, query_length: int) -> float:
    """Calculate query coverage from aligned query span."""
    if query_length <= 0:
        return 0.0
    aligned_query_span = abs(hit.query_end - hit.query_start) + 1
    return aligned_query_span / query_length


def _find_kasp_snp_query_positions(
    fam_seq: str,
    hex_seq: str,
) -> tuple[int, int, str, str] | None:
    """Find the 3' discriminatory SNP position in FAM/HEX primers.

    KASP allele-specific primers often carry extra 5' mismatches or length
    differences to tune assay behavior. The biological SNP is the 3'-most
    mismatch between the two allele-specific primers, not the first 5'
    mismatch.

    Returns:
        Tuple of (fam_query_pos_1based, hex_query_pos_1based, fam_allele,
        hex_allele), or ``None`` if the two sequences are identical over
        their 3' overlap.
    """
    if not fam_seq or not hex_seq:
        return None

    max_overlap = min(len(fam_seq), len(hex_seq))
    for offset_from_3prime in range(1, max_overlap + 1):
        fam_base = fam_seq[-offset_from_3prime]
        hex_base = hex_seq[-offset_from_3prime]
        if fam_base != hex_base:
            fam_query_pos = len(fam_seq) - offset_from_3prime + 1
            hex_query_pos = len(hex_seq) - offset_from_3prime + 1
            return fam_query_pos, hex_query_pos, fam_base, hex_base
    return None


def analyze_format1_loci(
    primer: KaspPrimer,
    results: list[BlastResult],
    min_identity: float = 95.0,
    min_coverage: float = 0.9,
    max_snp_distance: int = 10,
    max_amplicon_size: int = 1000,
    target_chroms: set[str] | None = None,
) -> list[ValidKaspLocus]:
    """Analyze valid Format1 loci with identity/coverage/chromosome/amplicon constraints.

    The SNP genomic coordinate is derived from the 3'-most mismatch between
    the two allele-specific primers. For every candidate pairing the SNP is
    mapped from query space to subject space independently for FAM and HEX,
    and pairings where the two coordinates disagree by more than
    ``max_snp_distance`` bp are discarded. ``target_chroms`` restricts BLAST
    hits to the listed chromosomes (used with an external ID→chromosome map).
    """
    fam_id = f"{primer.primer_id}_FAM"
    hex_id = f"{primer.primer_id}_HEX"
    common_id = f"{primer.primer_id}_Common"

    if not primer.fam_primer or not primer.hex_primer or not primer.common_primer:
        return []

    fam_query_length = len(primer.fam_primer)
    hex_query_length = len(primer.hex_primer)
    common_query_length = len(primer.common_primer)

    fam_hits = [
        r for r in results
        if r.query_id == fam_id
        and (not target_chroms or r.subject_id in target_chroms)
        and r.identity >= min_identity
        and _query_coverage(r, fam_query_length) >= min_coverage
    ]
    hex_hits = [
        r for r in results
        if r.query_id == hex_id
        and (not target_chroms or r.subject_id in target_chroms)
        and r.identity >= min_identity
        and _query_coverage(r, hex_query_length) >= min_coverage
    ]
    common_hits = [
        r for r in results
        if r.query_id == common_id
        and (not target_chroms or r.subject_id in target_chroms)
        and r.identity >= min_identity
        and _query_coverage(r, common_query_length) >= min_coverage
    ]

    if not fam_hits or not hex_hits or not common_hits:
        return []

    snp_info = _find_kasp_snp_query_positions(primer.fam_primer, primer.hex_primer)
    if not snp_info:
        return []
    fam_snp_query_pos, hex_snp_query_pos, ref_allele, alt_allele = snp_info

    valid_loci: list[ValidKaspLocus] = []
    chroms = {hit.subject_id for hit in fam_hits + hex_hits + common_hits}

    for chrom in sorted(chroms):
        cf = [h for h in fam_hits if h.subject_id == chrom]
        ch = [h for h in hex_hits if h.subject_id == chrom]
        cc = [h for h in common_hits if h.subject_id == chrom]
        if not cf or not ch or not cc:
            continue

        for fam_hit in cf:
            fam_center = (fam_hit.subject_start + fam_hit.subject_end) // 2
            fam_snp_pos = _query_to_subject_pos(fam_hit, fam_snp_query_pos)
            if fam_snp_pos is None:
                continue

            for hex_hit in ch:
                hex_center = (hex_hit.subject_start + hex_hit.subject_end) // 2
                hex_snp_pos = _query_to_subject_pos(hex_hit, hex_snp_query_pos)
                if hex_snp_pos is None:
                    continue

                # Use center-based pairing (like original pipeline)
                if abs(fam_center - hex_center) > max_snp_distance:
                    continue

                # The allele-specific SNP should map to nearly the same genomic base.
                if abs(fam_snp_pos - hex_snp_pos) > max_snp_distance:
                    continue

                snp_pos = fam_snp_pos

                for common_hit in cc:
                    all_pos = [
                        fam_hit.subject_start,
                        fam_hit.subject_end,
                        hex_hit.subject_start,
                        hex_hit.subject_end,
                        common_hit.subject_start,
                        common_hit.subject_end,
                    ]
                    amplicon_size = max(all_pos) - min(all_pos)
                    if amplicon_size >= max_amplicon_size:
                        continue

                    fam_coverage = _query_coverage(fam_hit, fam_query_length)
                    hex_coverage = _query_coverage(hex_hit, hex_query_length)
                    common_coverage = _query_coverage(common_hit, common_query_length)

                    query_length = common_query_length
                    alignment_length = common_hit.alignment_length
                    coverage = common_coverage

                    valid_loci.append(
                        ValidKaspLocus(
                            chrom=chrom,
                            snp_pos=snp_pos,
                            ref_allele=ref_allele,
                            alt_allele=alt_allele,
                            fam_pos=_format_interval(fam_hit.subject_start, fam_hit.subject_end),
                            hex_pos=_format_interval(hex_hit.subject_start, hex_hit.subject_end),
                            common_pos=_format_interval(common_hit.subject_start, common_hit.subject_end),
                            amplicon_size=amplicon_size,
                            query_length=query_length,
                            alignment_length=alignment_length,
                            coverage=coverage,
                            fam_coverage=fam_coverage,
                            hex_coverage=hex_coverage,
                            common_coverage=common_coverage,
                            fam_identity=fam_hit.identity,
                            hex_identity=hex_hit.identity,
                            common_identity=common_hit.identity,
                        )
                    )

    return valid_loci


def analyze_format2_loci(
    primer: KaspPrimer,
    results: list[BlastResult],
    min_identity: float = 95.0,
    min_coverage: float = 0.8,
    target_chroms: set[str] | None = None,
) -> tuple[list[ValidKaspLocus], str]:
    """Analyze valid Format2 loci with identity>=95% and coverage>=80%."""
    if not primer.flank_seq or primer.snp_pos is None:
        return [], "缺少flank序列或SNP位置信息"

    qlen = len(primer.flank_seq)
    if qlen == 0:
        return [], "flank序列长度为0"

    raw_hits = [r for r in results if r.query_id == primer.primer_id]
    if not raw_hits:
        return [], "无BLAST匹配"
    if target_chroms:
        raw_hits = [r for r in raw_hits if r.subject_id in target_chroms]
        if not raw_hits:
            return [], f"指定染色体({','.join(sorted(target_chroms))})无BLAST匹配"

    identity_hits = [r for r in raw_hits if r.identity >= min_identity]
    if not identity_hits:
        return [], f"Identity不足(<{min_identity}%)"

    passing_hits = [r for r in identity_hits if _query_coverage(r, qlen) >= min_coverage]
    if not passing_hits:
        max_cov = max(_query_coverage(r, qlen) for r in identity_hits)
        return [], f"Coverage不足(<{min_coverage:.0%})，当前最大{max_cov:.1%}"

    loci: list[ValidKaspLocus] = []
    for hit in passing_hits:
        snp_pos = _query_to_subject_pos(hit, primer.snp_pos)
        if snp_pos is None:
            continue

        cov = _query_coverage(hit, qlen)
        interval = _format_interval(hit.subject_start, hit.subject_end)
        loci.append(
            ValidKaspLocus(
                chrom=hit.subject_id,
                snp_pos=snp_pos,
                ref_allele=primer.ref_allele or "N",
                alt_allele=primer.alt_allele or "N",
                fam_pos=interval,
                hex_pos=interval,
                common_pos=interval,
                amplicon_size=interval[1] - interval[0],
                query_length=qlen,
                alignment_length=hit.alignment_length,
                coverage=cov,
                fam_coverage=cov,
                hex_coverage=cov,
                common_coverage=cov,
                fam_identity=hit.identity,
                hex_identity=hit.identity,
                common_identity=hit.identity,
            )
        )

    if not loci:
        return [], "SNP位点未落入比对区域"
    return loci, "成功"


def _hit_strand(hit: BlastResult) -> str:
    """Return '+' if subject is on the forward strand, '-' otherwise."""
    return "-" if hit.subject_start > hit.subject_end else "+"


def analyze_ssr_loci(
    primer: KaspPrimer,
    results: list[BlastResult],
    min_identity: float = 95.0,
    min_coverage: float = 0.9,
    max_amplicon_size: int = 1000,
    target_chroms: set[str] | None = None,
) -> tuple[list[ValidSsrLocus], str]:
    """Analyze valid SSR loci with identity/coverage/strand/amplicon constraints."""
    if not primer.forward_primer or not primer.reverse_primer:
        return [], "缺少Forward或Reverse引物序列"

    forward_id = f"{primer.primer_id}_F"
    reverse_id = f"{primer.primer_id}_R"
    forward_qlen = len(primer.forward_primer)
    reverse_qlen = len(primer.reverse_primer)

    forward_raw = [r for r in results if r.query_id == forward_id]
    reverse_raw = [r for r in results if r.query_id == reverse_id]
    if target_chroms:
        forward_raw = [r for r in forward_raw if r.subject_id in target_chroms]
        reverse_raw = [r for r in reverse_raw if r.subject_id in target_chroms]
    if not forward_raw or not reverse_raw:
        if target_chroms:
            return [], f"指定染色体({','.join(sorted(target_chroms))})无BLAST匹配"
        return [], "无BLAST匹配"

    forward_by_identity = [r for r in forward_raw if r.identity >= min_identity]
    reverse_by_identity = [r for r in reverse_raw if r.identity >= min_identity]
    if not forward_by_identity or not reverse_by_identity:
        return [], f"Identity不足(<{min_identity}%)"

    forward_hits = [
        r for r in forward_by_identity if _query_coverage(r, forward_qlen) >= min_coverage
    ]
    reverse_hits = [
        r for r in reverse_by_identity if _query_coverage(r, reverse_qlen) >= min_coverage
    ]
    if not forward_hits or not reverse_hits:
        return [], f"Coverage不足(<{min_coverage:.0%})"

    loci: list[ValidSsrLocus] = []
    chroms = {hit.subject_id for hit in forward_hits + reverse_hits}
    for chrom in sorted(chroms):
        f_hits = [h for h in forward_hits if h.subject_id == chrom]
        r_hits = [h for h in reverse_hits if h.subject_id == chrom]
        if not f_hits or not r_hits:
            continue

        for f_hit in f_hits:
            f_strand = _hit_strand(f_hit)
            for r_hit in r_hits:
                r_strand = _hit_strand(r_hit)
                if f_strand == r_strand:
                    continue
                positions = [
                    f_hit.subject_start,
                    f_hit.subject_end,
                    r_hit.subject_start,
                    r_hit.subject_end,
                ]
                amplicon_size = max(positions) - min(positions)
                if amplicon_size >= max_amplicon_size:
                    continue

                loci.append(
                    ValidSsrLocus(
                        chrom=chrom,
                        forward_pos=_format_interval(f_hit.subject_start, f_hit.subject_end),
                        reverse_pos=_format_interval(r_hit.subject_start, r_hit.subject_end),
                        forward_strand=f_strand,
                        reverse_strand=r_strand,
                        amplicon_size=amplicon_size,
                        forward_identity=f_hit.identity,
                        reverse_identity=r_hit.identity,
                        forward_coverage=_query_coverage(f_hit, forward_qlen),
                        reverse_coverage=_query_coverage(r_hit, reverse_qlen),
                    )
                )

    if not loci:
        return [], "无合法配对(同染色体/异链/扩增子长度)"
    return loci, "成功"


def _set_header_style(cell: Any) -> None:
    cell.fill = PatternFill("solid", fgColor="4F81BD")
    cell.font = Font(color="FFFFFF", bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _apply_fail_style(row_cells: list[Any]) -> None:
    for cell in row_cells:
        cell.fill = PatternFill("solid", fgColor="FFC7CE")
        cell.font = Font(color="9C0006")


def _apply_best_style(row_cells: list[Any]) -> None:
    for cell in row_cells:
        cell.fill = PatternFill("solid", fgColor="C6EFCE")
        cell.font = Font(color="006100", bold=True)


def create_excel_report(
    data: dict[str, dict[str, Any]],
    format_type: str,
    output_excel: Path,
    format1_min_coverage: float = 0.9,
    format2_min_coverage: float = 0.8,
    ssr_min_coverage: float = 0.9,
    ssr_max_amplicon: int = 1000,
    blast_threads: int = 1,
) -> None:
    """Generate 4-sheet KASP/SSR analysis report with formatting."""
    wb = Workbook()
    default_sheet = wb.active
    if default_sheet is not None:
        wb.remove(default_sheet)

    # Sheet 1: Methods
    ws1 = wb.create_sheet("分析方法说明")
    title = "SSR引物定位分析方法" if format_type == "ssr" else "KASP引物定位分析方法"
    ws1["A1"] = title
    ws1["A1"].font = Font(size=14, bold=True)
    ws1["A3"] = f"输入格式: {format_type}"
    ws1["A4"] = (
        "BLAST参数: blastn-short, evalue=1000, word_size=11, "
        f"num_threads={blast_threads}, dust=no, outfmt=6 std qlen"
    )
    if format_type == "format1":
        ws1["A6"] = "Format1判定规则:"
        ws1["A7"] = "1) FAM/HEX/Common Identity >= 95%"
        ws1["A8"] = f"2) FAM/HEX/Common Coverage >= {format1_min_coverage:.0%}"
        ws1["A9"] = "3) 三条引物必须在同一染色体"
        ws1["A10"] = "4) FAM和HEX SNP位置距离 <= 10bp"
        ws1["A11"] = "5) PCR扩增子长度 < 1000bp"
    elif format_type == "ssr":
        ws1["A6"] = "SSR判定规则:"
        ws1["A7"] = "1) Forward/Reverse Identity >= 95%"
        ws1["A8"] = f"2) Forward/Reverse Coverage >= {ssr_min_coverage:.0%}"
        ws1["A9"] = "3) Forward和Reverse必须在同一染色体上且处于互补链"
        ws1["A10"] = f"4) PCR扩增子长度 < {ssr_max_amplicon}bp"
    else:
        ws1["A6"] = "Format2判定规则:"
        ws1["A7"] = "1) Identity >= 95%"
        ws1["A8"] = f"2) Coverage >= {format2_min_coverage:.0%}"
        ws1["A9"] = "3) SNP位置必须落在比对覆盖区域"
    ws1.column_dimensions["A"].width = 120

    # Sheet 2: Valid loci
    sheet2_title = "有效SSR位点" if format_type == "ssr" else "有效KASP位点"
    ws2 = wb.create_sheet(sheet2_title)
    if format_type == "format1":
        headers = [
            "引物ID", "染色体", "SNP位置", "等位基因",
            "FAM位置", "HEX位置", "Common位置", "扩增子长度",
            "FAM Identity", "HEX Identity", "Common Identity",
            "FAM Coverage", "HEX Coverage", "Common Coverage",
        ]
    elif format_type == "ssr":
        headers = [
            "引物ID", "染色体",
            "Forward位置", "Forward链",
            "Reverse位置", "Reverse链",
            "扩增子长度",
            "Forward Identity", "Reverse Identity",
            "Forward Coverage", "Reverse Coverage",
            "状态",
        ]
    else:
        headers = [
            "引物ID", "染色体", "SNP位置", "等位基因",
            "比对区间", "Query长度", "比对长度", "Coverage",
            "Identity", "状态",
        ]
    ws2.append(headers)
    for cell in ws2[1]:
        _set_header_style(cell)

    for primer_id, entry in data.items():
        loci = entry.get("loci", [])
        fail_reason = entry.get("fail_reason", "")
        if not loci:
            row = [primer_id] + [""] * (len(headers) - 2) + [fail_reason or "无有效位点"]
            ws2.append(row)
            _apply_fail_style(list(ws2[ws2.max_row]))
            continue

        for loc in loci:
            if format_type == "format1":
                ws2.append([
                    primer_id, loc.chrom, loc.snp_pos, f"{loc.ref_allele}/{loc.alt_allele}",
                    f"{loc.fam_pos[0]}-{loc.fam_pos[1]}",
                    f"{loc.hex_pos[0]}-{loc.hex_pos[1]}",
                    f"{loc.common_pos[0]}-{loc.common_pos[1]}",
                    loc.amplicon_size, loc.fam_identity, loc.hex_identity, loc.common_identity,
                    loc.fam_coverage, loc.hex_coverage, loc.common_coverage,
                ])
            elif format_type == "ssr":
                ws2.append([
                    primer_id, loc.chrom,
                    f"{loc.forward_pos[0]}-{loc.forward_pos[1]}", loc.forward_strand,
                    f"{loc.reverse_pos[0]}-{loc.reverse_pos[1]}", loc.reverse_strand,
                    loc.amplicon_size,
                    loc.forward_identity, loc.reverse_identity,
                    loc.forward_coverage, loc.reverse_coverage,
                    "通过",
                ])
            else:
                ws2.append([
                    primer_id, loc.chrom, loc.snp_pos, f"{loc.ref_allele}/{loc.alt_allele}",
                    f"{loc.common_pos[0]}-{loc.common_pos[1]}",
                    loc.query_length, loc.alignment_length, loc.coverage, loc.common_identity, "通过",
                ])

    # Sheet 3: Best loci summary
    ws3 = wb.create_sheet("最可能位点汇总")
    if format_type == "ssr":
        summary_headers = ["引物ID", "染色体", "Forward位置", "Reverse位置", "扩增子长度", "评分", "备注"]
    else:
        summary_headers = ["引物ID", "染色体", "SNP位置", "等位基因", "评分", "备注"]
    ws3.append(summary_headers)
    for cell in ws3[1]:
        _set_header_style(cell)

    for primer_id, entry in data.items():
        loci = entry.get("loci", [])
        if not loci:
            fail_cells = [primer_id] + [""] * (len(summary_headers) - 3) + [
                0, entry.get("fail_reason", "无有效位点"),
            ]
            ws3.append(fail_cells)
            _apply_fail_style(list(ws3[ws3.max_row]))
            continue

        if format_type == "format1":
            def score(loc: ValidKaspLocus) -> float:
                identity_avg = (loc.fam_identity + loc.hex_identity + loc.common_identity) / 3.0
                return identity_avg - (loc.amplicon_size / 1000)
        elif format_type == "ssr":
            def score(loc: ValidSsrLocus) -> float:
                identity_avg = (loc.forward_identity + loc.reverse_identity) / 2.0
                penalty = loc.amplicon_size / ssr_max_amplicon if ssr_max_amplicon else 0.0
                return identity_avg - penalty
        else:
            def score(loc: ValidKaspLocus) -> float:
                return loc.common_identity * 0.7 + (loc.coverage * 100.0) * 0.3

        best_loc = max(loci, key=score)
        if format_type == "ssr":
            ws3.append([
                primer_id, best_loc.chrom,
                f"{best_loc.forward_pos[0]}-{best_loc.forward_pos[1]}",
                f"{best_loc.reverse_pos[0]}-{best_loc.reverse_pos[1]}",
                best_loc.amplicon_size,
                round(score(best_loc), 3), "最可能位点",
            ])
        else:
            ws3.append([
                primer_id, best_loc.chrom, best_loc.snp_pos,
                f"{best_loc.ref_allele}/{best_loc.alt_allele}",
                round(score(best_loc), 3), "最可能位点",
            ])
        _apply_best_style(list(ws3[ws3.max_row]))

    # Sheet 4: Statistics
    ws4 = wb.create_sheet("统计信息")
    total = len(data)
    passed = sum(1 for item in data.values() if item.get("loci"))
    failed = total - passed
    ws4.append(["指标", "数值"])
    ws4.append(["输入类型", format_type])
    ws4.append(["总引物数", total])
    ws4.append(["有效位点引物数", passed])
    ws4.append(["失败引物数", failed])
    ws4.append(["通过率", f"{(passed / total):.2%}" if total else "0.00%"])
    _set_header_style(ws4["A1"])
    _set_header_style(ws4["B1"])

    for sheet in [ws2, ws3, ws4]:
        for col_idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[chr(64 + col_idx)].width = 18

    wb.save(output_excel)
    logger.info(f"Excel report saved: {output_excel}")


@app.command()
def main(
    input: Annotated[
        Path,
        typer.Argument(
            help="Input TSV file (Format1, Format2 or SSR)",
            exists=True,
            readable=True,
            dir_okay=False,
        ),
    ],
    output: Annotated[
        Path,
        typer.Option(
            "--output",
            "-o",
            help="Output Excel report file",
        ),
    ] = Path("KASP引物定位结果.xlsx"),
    db_path: Annotated[
        Path,
        typer.Option(
            "--db",
            "-d",
            help="BLAST database path",
        ),
    ] = Path("blast_db/genome.fa"),
    id_chr_map: Annotated[
        Path | None,
        typer.Option(
            "--id-chr-map",
            help="Optional two-column file mapping primer ID to target chromosome(s)",
            exists=True,
            readable=True,
            dir_okay=False,
        ),
    ] = None,
    min_identity: Annotated[
        float,
        typer.Option(
            "--min-identity",
            "-i",
            help="Minimum BLAST identity threshold (%)",
        ),
    ] = 95.0,
    min_coverage_format1: Annotated[
        float,
        typer.Option(
            "--min-coverage-format1",
            help="Minimum coverage threshold for Format1 primers (0-1)",
        ),
    ] = 0.9,
    min_coverage: Annotated[
        float,
        typer.Option(
            "--min-coverage",
            "-c",
            help="Minimum coverage threshold for Format2 (0-1)",
        ),
    ] = 0.8,
    min_coverage_ssr: Annotated[
        float,
        typer.Option(
            "--min-coverage-ssr",
            help="Minimum coverage threshold for SSR primers (0-1)",
        ),
    ] = 0.9,
    max_amplicon: Annotated[
        int,
        typer.Option(
            "--max-amplicon",
            "-a",
            help="Maximum amplicon size for Format1",
        ),
    ] = 1000,
    max_amplicon_ssr: Annotated[
        int,
        typer.Option(
            "--max-amplicon-ssr",
            help="Maximum amplicon size for SSR (bp)",
        ),
    ] = 1000,
    threads: Annotated[
        int,
        typer.Option(
            "--threads",
            "-t",
            min=1,
            help="Number of threads passed to BLAST",
        ),
    ] = 1,
    keep_temp: Annotated[
        bool,
        typer.Option(
            "--keep-temp",
            "-k",
            help="Keep temporary files (query FASTA and BLAST results)",
        ),
    ] = False,
    verbose: Annotated[
        bool,
        typer.Option(
            "--verbose",
            "-v",
            help="Enable verbose logging",
        ),
    ] = False,
) -> None:
    """
    KASP/SSR primer mapping and reporting for Format1/Format2/SSR TSV inputs.

    Format1: ID\tFAM\tHEX\tCommon (with dye tags)
    Format2: ID\tFlank (with [REF/ALT] SNP marker)
    SSR:     ID\tForward\tReverse (no dye tag, no SNP marker)
    """
    if verbose:
        logger.remove()
        logger.add(
            sys.stderr,
            format="<green>{time:HH:mm:ss}</green> | <level>{level: <8}</level> | <level>{message}</level>",
            level="DEBUG",
            colorize=True,
        )

    # Temp files
    query_file = Path("kasp_query.fa")
    blast_output = Path("kasp_blast_results.txt")

    try:
        with Progress(
            SpinnerColumn(),
            TextColumn("[progress.description]{task.description}"),
            console=console,
            transient=True,
        ) as progress:
            # Detect format
            task = progress.add_task("Detecting format...", total=None)
            format_type = detect_format(input)
            logger.info(f"Detected format: {format_type}")
            progress.update(task, description=f"Format: {format_type}")

            # Parse input
            progress.update(task, description="Parsing input...")
            if format_type == "format1":
                primers = parse_kasp1(input)
            elif format_type == "ssr":
                primers = parse_ssr(input)
            else:
                primers = parse_kasp2(input)
            progress.update(task, description=f"Parsed {len(primers)} primers")

            # Load optional ID-to-chromosome hints
            id_chroms = load_id_chr_map(id_chr_map)
            if id_chroms:
                progress.update(task, description=f"Loaded {len(id_chroms)} chromosome hints")

            # Create BLAST query
            progress.update(task, description="Creating BLAST query...")
            create_blast_query_file(primers, query_file)

            # Run BLAST
            progress.update(task, description="Running BLAST...")
            run_blast(query_file, db_path, blast_output, threads=threads)
            progress.update(task, description="BLAST completed")

            # Parse BLAST results
            progress.update(task, description="Parsing BLAST results...")
            blast_results = parse_blast_results(blast_output)
            progress.update(task, description=f"Parsed {len(blast_results)} hits")

            # Analyze loci
            progress.update(task, description="Analyzing loci...")
            report_data: dict[str, dict[str, Any]] = {}
            if format_type == "format1":
                for primer in primers:
                    primer_chroms = id_chroms.get(primer.primer_id)
                    loci = analyze_format1_loci(
                        primer, blast_results,
                        min_identity=min_identity,
                        min_coverage=min_coverage_format1,
                        max_amplicon_size=max_amplicon,
                        target_chroms=primer_chroms,
                    )
                    report_data[primer.primer_id] = {
                        "primer": primer,
                        "loci": loci,
                        "fail_reason": "" if loci else (
                            f"指定染色体({','.join(sorted(primer_chroms))})未通过Format1过滤条件"
                            if primer_chroms else "未通过Format1过滤条件"
                        ),
                    }
            elif format_type == "ssr":
                for primer in primers:
                    primer_chroms = id_chroms.get(primer.primer_id)
                    loci, reason = analyze_ssr_loci(
                        primer, blast_results,
                        min_identity=min_identity,
                        min_coverage=min_coverage_ssr,
                        max_amplicon_size=max_amplicon_ssr,
                        target_chroms=primer_chroms,
                    )
                    report_data[primer.primer_id] = {
                        "primer": primer,
                        "loci": loci,
                        "fail_reason": "" if loci else reason,
                    }
            else:
                for primer in primers:
                    primer_chroms = id_chroms.get(primer.primer_id)
                    loci, reason = analyze_format2_loci(
                        primer, blast_results,
                        min_identity=min_identity,
                        min_coverage=min_coverage,
                        target_chroms=primer_chroms,
                    )
                    report_data[primer.primer_id] = {
                        "primer": primer,
                        "loci": loci,
                        "fail_reason": "" if loci else reason,
                    }
            progress.update(task, description="Analysis completed")

            # Create report
            progress.update(task, description="Creating Excel report...")
            create_excel_report(
                report_data,
                format_type,
                output,
                format1_min_coverage=min_coverage_format1,
                format2_min_coverage=min_coverage,
                ssr_min_coverage=min_coverage_ssr,
                ssr_max_amplicon=max_amplicon_ssr,
                blast_threads=threads,
            )

        # Summary
        total = len(report_data)
        success = sum(1 for item in report_data.values() if item["loci"])

        summary_title = "SSR Mapping Summary" if format_type == "ssr" else "KASP Mapping Summary"
        summary = Table(title=summary_title)
        summary.add_column("Metric", style="cyan")
        summary.add_column("Value", style="green")
        summary.add_row("Input file", str(input))
        summary.add_row("Output file", str(output))
        summary.add_row("Format type", format_type)
        summary.add_row("Total primers", str(total))
        summary.add_row("Successful", str(success))
        summary.add_row("Failed", str(total - success))
        if id_chr_map is not None:
            summary.add_row("ID-chr map entries", str(len(id_chroms)))
        summary.add_row("Success rate", f"{(success / total):.1%}" if total else "N/A")
        console.print(summary)

        logger.success(f"Report saved: {output}")

    except FileNotFoundError as exc:
        logger.error(f"File not found: {exc}")
        raise typer.Exit(code=2) from exc
    except RuntimeError as exc:
        logger.error(f"Runtime error: {exc}")
        raise typer.Exit(code=3) from exc
    except ValueError as exc:
        logger.error(f"Validation error: {exc}")
        raise typer.Exit(code=4) from exc
    finally:
        # Cleanup temp files
        if not keep_temp:
            for f in [query_file, blast_output]:
                if f.exists():
                    f.unlink()
                    logger.debug(f"Removed temp file: {f}")
        else:
            logger.info(f"Kept temp files: {query_file}, {blast_output}")


if __name__ == "__main__":
    app()
